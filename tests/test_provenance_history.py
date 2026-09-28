from copy import deepcopy
import io

import pytest

from app.aggregate import aggregate
from app.export import _evidence, build_csv, build_markdown, build_pdf, build_pptx
from app.providers import c2pa_provider as provider
from app.schema import Verdict


def store(*actions):
    return {"active_manifest": "active", "validation_state": "Trusted",
            "validation_results": {"activeManifest": {"success": [], "failure": []}},
            "manifests": {"active": {"claim_generator": "Test exporter",
                "assertions": [{"label": "c2pa.actions.v2", "data": {"actions": list(actions)}}]}}}


def test_transcoded_creation_history_is_visible_but_not_a_synthetic_declaration():
    data = store({"action": "c2pa.opened"}, {"action": "c2pa.transcoded"})
    data["manifests"]["active"]["ingredients"] = [{"relationship": "parentOf", "active_manifest": "parent"}]
    data["manifests"]["parent"] = {"claim_generator": "Test creation app", "assertions": [
        {"label": "c2pa.actions", "data": {"actions": [{"action": "c2pa.created"}]}}]}
    row = provider._interpret(data, 0, policy_approved=True)
    assert row.verdict == Verdict.INCONCLUSIVE
    assert row.rating is None
    assert row.raw["history_complete"]
    assert row.raw["history"][1]["generator"] == "Test creation app"
    assert row.raw["history"][1]["actions"] == ["c2pa.created"]
    assert "No recognised" in row.summary


def test_ai_ingredient_does_not_decide_whole_file_even_with_trusted_active_record():
    data = store({"action": "c2pa.placed"})
    data["manifests"]["active"]["ingredients"] = [{"relationship": "componentOf", "active_manifest": "part"}]
    data["manifests"]["part"] = store({"digitalSourceType": provider._AI_SOURCE})["manifests"]["active"]
    row = provider._interpret(data, 0, policy_approved=True)
    assert row.raw["history"][1]["source_types"] == [provider._AI_SOURCE]
    assert row.verdict == Verdict.INCONCLUSIVE
    assert aggregate([row], "", "general")["overall_verdict"] == "inconclusive"


@pytest.mark.parametrize("label", ["c2pa.actions", "c2pa.actions.v2", "c2pa.actions.v2__2"])
def test_templates_and_multiple_assertions_resolve_without_mutating_input(label):
    data = store({"action": "c2pa.created"})
    assertion = data["manifests"]["active"]["assertions"][0]
    assertion["label"] = label
    assertion["data"]["templates"] = [
        {"action": "c2pa.created", "digitalSourceType": provider._AI_SOURCE},
        {"action": "*", "digitalSourceType": provider._CAMERA_SOURCE}]
    original = deepcopy(data)
    row = provider._interpret(data, 0, policy_approved=True)
    assert row.verdict == Verdict.SYNTHETIC_LIKELY
    assert data == original
    assertion["data"]["actions"][0]["digitalSourceType"] = provider._CAMERA_SOURCE
    assert provider._interpret(data, 0, policy_approved=True).verdict == Verdict.AUTHENTIC_LIKELY


@pytest.mark.parametrize("label", ["c2pa.actions.v3", "c2pa.actions.v2__oops", "other.c2pa.actions"])
def test_unknown_assertion_labels_do_not_create_overrides(label):
    data = store({"digitalSourceType": provider._AI_SOURCE})
    data["manifests"]["active"]["assertions"][0]["label"] = label
    assert provider._interpret(data, 0, policy_approved=True).verdict == Verdict.INCONCLUSIVE


def test_ai_disclosure_is_reported_as_context_without_inventing_whole_file_generation():
    data = store({"action": "c2pa.edited"})
    data["manifests"]["active"]["assertions"].append({"label": "c2pa.ai-disclosure__1", "data": {
        "modelType": "diffusion", "modelName": "Fixture model",
        "contentProfile": {"humanOversightLevel": "human_validated"}}})
    row = provider._interpret(data, 0, policy_approved=True)
    assert row.raw["history"][0]["ai_disclosures"][0]["modelName"] == "Fixture model"
    assert row.verdict == Verdict.INCONCLUSIVE
    assert "whole-file origin assessed separately" in " ".join(row.evidence)


@pytest.mark.parametrize("scope", ["region", "assertion", "related"])
def test_scoped_ai_declarations_never_become_whole_file_verdicts(scope):
    action = {"action": "c2pa.edited", "digitalSourceType": provider._AI_SOURCE}
    data = store(action)
    assertion = data["manifests"]["active"]["assertions"][0]["data"]
    if scope == "region":
        action["changes"] = [{"region": [{"type": "temporal"}]}]
    elif scope == "assertion":
        assertion["metadata"] = {"regionOfInterest": {"region": []}}
    else:
        assertion["actions"] = [{"action": "c2pa.edited", "related": [action]}]
    row = provider._interpret(data, 0, policy_approved=True)
    assert row.verdict == Verdict.INCONCLUSIVE
    assert row.raw["scoped_source_types"] == [provider._AI_SOURCE]
    assert "Whole-file coverage is unresolved" in row.summary


@pytest.mark.parametrize("value", [True, False, None])
def test_action_completeness_is_only_a_signer_declaration(value):
    data = store({"action": "c2pa.opened"})
    if value is not None:
        data["manifests"]["active"]["assertions"][0]["data"]["allActionsIncluded"] = value
    row = provider._interpret(data, 0)
    assert row.raw["history"][0]["all_actions_included"] is value
    assert row.verdict == Verdict.INCONCLUSIVE


def test_scoped_ai_edit_does_not_leave_a_camera_only_origin_verdict():
    data = store({"action": "c2pa.created", "digitalSourceType": provider._CAMERA_SOURCE},
                 {"action": "c2pa.edited", "digitalSourceType": provider._AI_SOURCE, "changes": []})
    row = provider._interpret(data, 0, policy_approved=True)
    assert row.verdict == Verdict.INCONCLUSIVE
    assert "coexist" in row.summary


def test_history_only_follows_embedded_reachable_ingredients_and_handles_cycles():
    data = store({"action": "c2pa.opened"})
    data["manifests"]["unrelated"] = {"claim_generator": "Do not display"}
    data["manifests"]["active"]["ingredients"] = [{"active_manifest": "active"}]
    row = provider._interpret(data, 0)
    assert len(row.raw["history"]) == 1
    assert row.raw["history_complete"]
    data["manifests"]["active"]["ingredients"].append({"active_manifest": "missing"})
    assert not provider._interpret(data, 0).raw["history_complete"]


def test_provenance_exports_keep_summary_and_evidence_beyond_five_lines():
    row = provider._interpret(store({"action": "c2pa.opened"}), 0).to_dict()
    row["evidence"] += [f"Fixture observation {i}" for i in range(10)] + ["Final trust caveat"]
    report = {"providers": [row], "frames": [], "consensus": aggregate([], "", "general"),
              "meta": {"tool": "SDA Vision", "version": "test", "filename": "fixture.mp4",
                       "kind": "video", "generated_at": "2026-09-24T00:00:00Z", "mode": "general",
                       "frame_count": 0, "models": [], "notes": [], "elapsed_ms": 0}}
    assert row["summary"] in _evidence(row)
    assert "Final trust caveat" in build_markdown(report)
    assert "Final trust caveat" in build_csv(report)
    import io
    from pypdf import PdfReader
    pdf = PdfReader(io.BytesIO(build_pdf(report)))
    assert "Final trust caveat" in "".join(page.extract_text() for page in pdf.pages)
    from pptx import Presentation
    deck = Presentation(io.BytesIO(build_pptx(report)))
    # Research reports now put evidence in structured tables as well as text frames.
    text = "\n".join(shape.text if shape.has_text_frame else "\n".join(
        cell.text for row in shape.table.rows for cell in row.cells) if shape.has_table else ""
        for slide in deck.slides for shape in slide.shapes)
    assert "Final trust caveat" in text
