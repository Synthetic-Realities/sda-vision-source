import io
import re

import pytest
from pptx import Presentation

from app.aggregate import aggregate
from app.export import build_pptx
from scripts import package_candidate as candidate


def test_candidate_contains_local_assets_imported_by_frontend():
    """A source archive must retain assets needed to rebuild its frontend."""
    paths = set(candidate.candidate_files())
    assert not {"AGENTS.md", "CONTRIBUTING.md"} & {p.name for p in paths}
    for source in paths:
        if source.suffix not in {".ts", ".tsx"}:
            continue
        for asset in re.findall(r'from\s+[\"\']([^\"\']+\.(?:png|jpe?g|webp|svg|ttf|woff2?)(?:\?url)?)[\"\']',
                                source.read_text()):
            if asset.startswith("."):
                target = (source.parent / asset.split("?")[0]).resolve()
                assert target in paths, f"Missing packaged frontend asset: {target.relative_to(candidate.ROOT)}"


def test_long_pptx_details_are_paginated_and_preserved():
    report = {
        "meta": {"tool": "SDA Vision", "filename": "long-fixture.png", "generated_at": "2026-09-19",
                 "models": [], "notes": [], "kind": "image", "frame_count": 1},
        "providers": [], "consensus": aggregate([], "", "general"), "frames": [],
        "second_opinion_gemini": "Researcher note, not verified detection. " * 180 + "UNIQUE END MARKER",
    }
    deck = Presentation(io.BytesIO(build_pptx(report)))
    assert len(deck.slides) > 3
    text = "\n".join(cell.text for slide in deck.slides for shape in slide.shapes if shape.has_table
                     for row in shape.table.rows for cell in row.cells)
    assert "UNIQUE END MARKER" in " ".join(text.split())
    for slide in deck.slides:
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= deck.slide_width
            assert shape.top + shape.height <= deck.slide_height
    for slide in deck.slides:
        assert any(shape.has_table for shape in slide.shapes)
        for shape in slide.shapes:
            if shape.has_table:
                assert all(len(cell.text.splitlines()) <= 21 for row in shape.table.rows for cell in row.cells)


def test_candidate_allowlist_excludes_private_files(monkeypatch, tmp_path):
    monkeypatch.setattr(candidate, "FILES", {"README.md"})
    monkeypatch.setattr(candidate, "TREES", {"app": {".py"}})
    (tmp_path / "README.md").write_text("source")
    (tmp_path / "app").mkdir()
    (tmp_path / "app/main.py").write_text("source")
    for folder in (".git", "dev_runs", "labelled-corpus", "trust-store", ".review-cache"):
        (tmp_path / folder).mkdir()
        (tmp_path / folder / "private.py").write_text("DO NOT PACKAGE")
    (tmp_path / ".env").write_text("SECRET=DO NOT PACKAGE")
    (tmp_path / "examples").mkdir()
    (tmp_path / "examples/personal-film.mp4").write_bytes(b"PRIVATE MEDIA")
    assert {p.relative_to(tmp_path).as_posix() for p in candidate.candidate_files(tmp_path)} == {
        "README.md", "app/main.py"}


def test_candidate_rejects_symlinked_sources(monkeypatch, tmp_path):
    monkeypatch.setattr(candidate, "FILES", {"README.md"})
    monkeypatch.setattr(candidate, "TREES", {})
    (tmp_path / "private.md").write_text("DO NOT PACKAGE")
    (tmp_path / "README.md").symlink_to(tmp_path / "private.md")
    with pytest.raises(ValueError, match="unsafe"):
        candidate.candidate_files(tmp_path)


def test_package_can_use_review_build_without_replacing_working_frontend(monkeypatch, tmp_path):
    import zipfile
    monkeypatch.setattr(candidate, "FILES", {"README.md"})
    monkeypatch.setattr(candidate, "TREES", {"frontend/dist": {".html", ".js"}})
    (tmp_path / "README.md").write_text("Source candidate")
    working = tmp_path / "frontend/dist"; working.mkdir(parents=True)
    (working / "index.html").write_text("Historical working frontend")
    review = tmp_path / ".review-cache/fresh-frontend"; review.mkdir(parents=True)
    (review / "index.html").write_text("Reviewed current frontend")
    (review / "app.js").write_text("Reviewed current JavaScript")
    (review / "guides").mkdir()
    (review / "guides/SDA_Vision_Facilitator_Guide_v1.6-review.pdf").write_bytes(b"Reviewed guide")
    output = candidate.package(tmp_path, review)
    with zipfile.ZipFile(output) as archive:
        assert archive.read("sda-vision-candidate/frontend/dist/index.html") == b"Reviewed current frontend"
        assert archive.read("sda-vision-candidate/frontend/dist/app.js") == b"Reviewed current JavaScript"
        assert archive.read("sda-vision-candidate/frontend/dist/guides/SDA_Vision_Facilitator_Guide_v1.6-review.pdf") == b"Reviewed guide"
        assert not any(".review-cache/" in name for name in archive.namelist())
    assert (working / "index.html").read_text() == "Historical working frontend"


def test_candidate_keeps_extensionless_notices_and_skips_directories(monkeypatch, tmp_path):
    monkeypatch.setattr(candidate, "FILES", set())
    monkeypatch.setattr(candidate, "TREES", {"third_party": {"", ".md"}})
    notices = tmp_path / "third_party" / "example-package"
    notices.mkdir(parents=True)
    (notices / "LICENSE").write_text("Exact third-party notice")
    (tmp_path / "AGENTS.md").write_text("Local collaboration guidance")
    (tmp_path / "CONTRIBUTING.md").write_text("Local contribution guidance")
    assert candidate.candidate_files(tmp_path) == [notices / "LICENSE"]
    (notices / "linked-folder").symlink_to(notices, target_is_directory=True)
    with pytest.raises(ValueError, match="unsafe"):
        candidate.candidate_files(tmp_path)


def test_current_guide_and_game_assets_are_packaged():
    paths = {p.relative_to(candidate.ROOT).as_posix() for p in candidate.candidate_files()}
    for name in candidate.GUIDE_REVIEW_ASSETS | candidate.GUIDE_RELEASE_ASSETS:
        assert "frontend/public/guides/" + name in paths
        assert ("docs/guides/" if name.endswith(".pdf") else "docs/images/") + name in paths
        assert "frontend/dist/guides/" + name in candidate.BUILD_FILES
    for path in (candidate.ROOT / "frontend/public/game").iterdir():
        assert path.relative_to(candidate.ROOT).as_posix() in paths
    for name in ("facilitator_workflow", "analysis_citation", "game_view"):
        assert f"tests/{name}.mjs" in paths


def test_review_build_keeps_current_pdfs_and_game_notice(monkeypatch, tmp_path):
    import zipfile
    monkeypatch.setattr(candidate, "FILES", {"README.md"})
    monkeypatch.setattr(candidate, "TREES", {"frontend/dist": {".html", ".js", ".png"}})
    (tmp_path / "README.md").write_text("Source candidate")
    review = tmp_path / ".review-cache/fresh-frontend"
    (review / "guides").mkdir(parents=True)
    (review / "game").mkdir()
    (review / "index.html").write_text("Reviewed frontend")
    for name in candidate.GUIDE_REVIEW_ASSETS:
        (review / "guides" / name).write_bytes(name.encode())
    (review / "game/LICENSE.txt").write_text("Game artwork notice")
    output = candidate.package(tmp_path, review)
    with zipfile.ZipFile(output) as archive:
        for name in candidate.GUIDE_REVIEW_ASSETS:
            assert archive.read("sda-vision-candidate/frontend/dist/guides/" + name) == name.encode()
        assert archive.read("sda-vision-candidate/frontend/dist/game/LICENSE.txt") == b"Game artwork notice"
