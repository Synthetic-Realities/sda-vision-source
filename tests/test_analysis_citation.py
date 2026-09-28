"""Software attribution in readable exports; original research records stay intact."""
import copy
import csv
import io
import json
from pathlib import Path

import pytest
from pypdf import PdfReader
from pptx import Presentation

from app.citation import software_citation
from app.export import build_csv, build_markdown, build_pdf, build_pptx

def sample():
    return json.loads((Path(__file__).parent / "fixtures/workshop-review/podcast.json").read_text())

def credited(text):
    text = " ".join(text.split())
    for phrase in ("Martin, S. (2026)", "Dr Sam Martin", "Manchester Metropolitan University",
                   "Smart Data Research UK, a UKRI investment; Grant number UKRI4010.",
                   "0000-0002-4466-8374", "Software creator"):
        assert phrase in text

@pytest.mark.parametrize("kind", ["image", "audio", "video", "pdf", "pptx", "text"])
def test_citations_retain_report_identity_and_do_not_change_findings(kind):
    report = sample()
    report["meta"]["kind"] = kind
    original = copy.deepcopy(report)
    for text in (build_markdown(report), build_csv(report)):
        credited(text)
        assert report["meta"]["filename"] in text
        assert report["meta"]["generated_at"] in text
    rows = list(csv.DictReader(io.StringIO(build_csv(report))))
    for row in rows:
        credited(row["record_note"])
    assert report == original

def test_pdf_and_powerpoint_keep_credit_text_with_audio_scope():
    report = sample()
    pdf = PdfReader(io.BytesIO(build_pdf(report)))
    credited(" ".join(page.extract_text() for page in pdf.pages))
    deck = Presentation(io.BytesIO(build_pptx(report)))
    text = " ".join(cell.text for slide in deck.slides for shape in slide.shapes if shape.has_table
                    for row in shape.table.rows for cell in row.cells)
    credited(text)
    assert "transcript" in text.lower()

def test_missing_version_is_not_invented():
    assert "version not recorded" in software_citation()
    assert "version 0.1.0" in software_citation("0.1.0")
