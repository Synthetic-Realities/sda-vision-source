import asyncio
import hashlib
import io
import json
from dataclasses import replace

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from app import main, pipeline
from app.aggregate import aggregate, aggregate_frames
from app.config import METHOD_VERSION
from app.export import _score_basis, build_csv, build_markdown
from app.ingest import ingest
from app.schema import Kind, ProviderResult


@pytest.fixture
def client(monkeypatch):
    monkeypatch.setattr(main, "settings", replace(main.settings, dev_mode=False))
    async def no_network(*args, **kwargs):
        raise AssertionError("Unexpected cloud-capable call")
    monkeypatch.setattr(main, "analyse_deep", no_network)
    monkeypatch.setattr(main, "score_opinion", no_network)
    return TestClient(main.app)


def report():
    row = ProviderResult("example", "Example model", Kind.VISION, rating=42)
    return {
        "meta": {"tool": "SDA Vision", "version": "test", "filename": "fixture.png",
                 "generated_at": "2026-09-19T00:00:00Z", "kind": "image",
                 "frame_count": 0, "frames_found": 0, "models": [], "notes": []},
        "providers": [row.to_dict()], "frames": [],
        "consensus": aggregate([row], "", "general"),
    }


@pytest.mark.parametrize("field,value", [
    ("frames", 1), ("frames", [None]), ("providers", 1), ("providers", [None]),
    ("providers", [{"id": "x", "kind": "unknown"}]), ("meta", None), ("consensus", None),
])
def test_malformed_reports_are_rejected_before_processing(client, field, value):
    r = report()
    r[field] = value
    assert client.post("/api/export/md", json={"report": r}).status_code == 400
    assert client.post("/api/analyse/deep", data={"prior_report": json.dumps(r)},
                       files={"file": ("fixture.png", b"local only", "image/png")}).status_code == 400


@pytest.mark.parametrize("payload", [{"report": None}, {"report": []}, {"report": {}}])
def test_invalid_single_exports(client, payload):
    assert client.post("/api/export/pdf", json=payload).status_code == 400


@pytest.mark.parametrize("items", [[None], [1], [{"report": None}], 1, [], {}])
def test_invalid_batch_exports(client, items):
    assert client.post("/api/export/batch/pdf", json={"items": items}).status_code == 400


def test_note_scoring_requires_explicit_consent(client):
    assert client.post("/api/opinion-score", json={"text": "example"}).status_code == 400
    assert client.post("/api/opinion-score",
                       json={"text": [], "consent_to_anthropic": True}).status_code == 400


@pytest.mark.parametrize("difference", ["file", "method", "context", "extraction"])
def test_deep_pass_rejects_wrong_identity_without_cloud_calls(client, difference):
    raw = b"synthetic fixture bytes"
    r = report()
    r["meta"].update({
        "input_sha256": hashlib.sha256(raw).hexdigest(),
        "method_version": METHOD_VERSION,
        "context_sha256": pipeline.context_hash("", "general"),
        "extraction_policy": pipeline.extraction_policy(),
    })
    if difference == "file":
        raw += b" changed"
    elif difference == "method":
        r["meta"]["method_version"] = "old"
    elif difference == "context":
        r["meta"]["context_sha256"] = pipeline.context_hash("changed", "general")
    else:
        r["meta"]["extraction_policy"] = {}
    assert client.post("/api/analyse/deep", data={"prior_report": json.dumps(r)},
                       files={"file": ("fixture.png", raw, "image/png")}).status_code == 400


def test_public_mode_allows_analysis_but_does_not_save(monkeypatch, client):
    async def fake_analyse(ing, caption, mode):
        return report()
    def forbidden_save(*args):
        raise AssertionError("Public mode must not save reports")
    monkeypatch.setattr(main, "analyse", fake_analyse)
    monkeypatch.setattr(main, "_save_run", forbidden_save)
    response = client.post("/api/analyse", files={"file": ("fixture.txt", b"local fixture")})
    assert response.status_code == 200
    assert client.get("/api/corpus").json()["available"] is False


def test_text_only_pptx_uses_text_route_and_records_file_identity(monkeypatch):
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[1]).shapes.title.text = "Local text-only fixture"
    output = io.BytesIO()
    deck.save(output)
    raw = output.getvalue()
    calls = []
    async def fake_text(text, caption, mode):
        calls.append(text)
        return report()
    monkeypatch.setattr(pipeline, "_text_report", fake_text)
    item = ingest(raw, "fixture.pptx")
    assert not item.frames
    result = asyncio.run(pipeline.analyse(item, "", "general"))
    assert calls and "Local text-only fixture" in calls[0]
    assert result["meta"]["input_sha256"] == hashlib.sha256(raw).hexdigest()
    assert result["meta"]["method_version"] == METHOD_VERSION
    assert any("Text-only analysis" in n for n in result["meta"]["notes"])


def test_score_basis_matches_actual_frame_aggregation():
    r = report()
    r["consensus"] = aggregate_frames([
        {"overall_verdict": "inconclusive", "overall_rating": 20},
        {"overall_verdict": "inconclusive", "overall_rating": 80},
    ], [], "", "general")
    assert r["consensus"]["overall_rating"] == 50
    assert "mean of available frame scores" in _score_basis(r)
    assert "mean of available frame scores" in build_markdown(r)
    assert "mean of available frame scores" in build_csv(r)


def test_old_provenance_scores_are_not_relabelled_as_model_medians():
    r = report()
    r["consensus"]["overall_rating"] = 95
    r["consensus"]["decision_trace"] = [{"step": "override", "outcome": "synthetic_likely"}]
    assert "legacy fixed" in _score_basis(r)


def test_malformed_per_frame_provider_details_are_rejected(client):
    r = report()
    r["providers"][0]["raw"] = {"per_frame": 1}
    assert client.post("/api/export/md", json={"report": r}).status_code == 400


@pytest.mark.parametrize("fmt", ["md", "csv", "pdf", "pptx"])
def test_valid_export_still_works(client, fmt):
    response = client.post("/api/export/" + fmt, json={"report": report()})
    assert response.status_code == 200
    assert response.content
