import asyncio
import copy
import hashlib
import io
import json
import shutil
import subprocess
from dataclasses import replace

import httpx
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from app import audio, main, pipeline
from app.aggregate import aggregate
from app.export import build_csv, build_markdown, build_pdf, build_pptx
from app.ingest import Ingested, ingest
from app.media import load_media
from app.providers import audio_assessment, c2pa_provider, suno_credentials
from app.request_validation import validate_report
from app.schema import Kind, ProviderResult, Status


def inspection():
    return {"method_version": audio.METHOD, "input_sha256": "a" * 64, "status": "present",
            "stream_count": 1, "tracks": [{"index": 1, "codec": "aac", "channels": 2,
                "sample_rate": 44100, "duration_seconds": 40.0, "decode_status": "decoded",
                "samples": [{"start_seconds": 0.0, "duration_seconds": 3.0, "peak_dbfs": -20.0,
                             "signal_above_minus_60_dbfs": True}]}],
            "synthetic_audio_assessed": False, "provenance_scope": "Original container; audio scope unknown.",
            "notes": [audio.LIMIT]}


@pytest.mark.skipif(not shutil.which("ffmpeg"), reason="ffmpeg unavailable")
@pytest.mark.parametrize("silent", [False, True])
def test_local_audio_signal_and_original_identity(tmp_path, silent):
    path = tmp_path / "not-a-platform-name.mp4"
    source = "anullsrc=r=44100:cl=stereo" if silent else "sine=frequency=440:sample_rate=44100"
    subprocess.run(["ffmpeg", "-v", "error", "-f", "lavfi", "-i", source,
                    "-t", "2", "-c:a", "aac", str(path)], check=True)
    raw = path.read_bytes()
    value = audio.inspect_audio(raw, path.name)
    assert value["status"] == "present" and value["stream_count"] == 1
    assert value["tracks"][0]["decode_status"] == "decoded"
    assert value["tracks"][0]["samples"][0]["signal_above_minus_60_dbfs"] is not silent
    assert value["input_sha256"] == hashlib.sha256(raw).hexdigest()
    assert path.read_bytes() == raw
    excerpt, seconds = audio.extract_excerpt(raw, path.name, 0, 0)
    assert excerpt.startswith(b"RIFF") and 1.5 <= seconds <= 2.1
    assert value["synthetic_audio_assessed"] is False
    with pytest.raises(ValueError, match="track not found"):
        audio.extract_excerpt(raw, path.name, 2, 0)


def test_absent_unknown_and_invalid_are_distinct(monkeypatch):
    monkeypatch.setattr(audio, "which", lambda _: None)
    assert audio.inspect_audio(b"bad", "test.mp4")["status"] == "unavailable"
    monkeypatch.setattr(audio, "which", lambda _: "available")
    monkeypatch.setattr(audio, "_probe", lambda _: {"streams": []})
    assert audio.inspect_audio(b"bad", "test.mp4")["status"] == "absent"
    def failed(_):
        raise subprocess.TimeoutExpired("ffprobe", 10)
    monkeypatch.setattr(audio, "_probe", failed)
    assert audio.inspect_audio(b"bad", "test.mp4")["status"] == "error"


def test_multiple_tracks_caps_and_decode_failure(monkeypatch):
    monkeypatch.setattr(audio, "which", lambda _: "available")
    monkeypatch.setattr(audio, "_probe", lambda _: {"streams": [
        {"codec_type": "audio", "index": i, "duration": "40"} for i in range(7)]})
    def failed(*args, **kwargs):
        raise ValueError("decode failed")
    monkeypatch.setattr(audio, "_decode", failed)
    value = audio.inspect_audio(b"raw", "unknown.mp4")
    assert value["status"] == "present" and value["stream_count"] == 7
    assert len(value["tracks"]) == 4
    assert all(t["decode_status"] == "error" for t in value["tracks"])
    assert audio._number("nan") is None and audio._number("inf") is None


def test_video_gets_local_audio_record_without_extra_cloud_calls(monkeypatch):
    base = {"consensus": aggregate([], "", "general"), "providers": [], "frames": []}
    async def frames(*_):
        return copy.deepcopy(base)
    monkeypatch.setattr(pipeline, "_analyse_frames", frames)
    monkeypatch.setattr(pipeline, "inspect_audio", lambda *_: inspection())
    ing = Ingested("video", "unlabelled.mp4", load_media(b"original", "unlabelled.mp4"), frames=[None])
    result = asyncio.run(pipeline.analyse(ing, "", "general"))
    assert result["audio_inspection"]["status"] == "present"
    assert result["consensus"] == base["consensus"]
    assert "audio_assessment" not in result and "suno_check" not in result


def test_video_without_frames_still_checks_original_credentials_and_audio(monkeypatch):
    events = []
    def original(media):
        assert media.raw_bytes == b"original"
        events.append("original")
        return ProviderResult("c2pa", "C2PA", Kind.PROVENANCE, Status.OK, summary="Original inspected")
    def soundtrack(*_):
        events.append("audio")
        return inspection()
    monkeypatch.setattr(c2pa_provider, "run", original)
    monkeypatch.setattr(pipeline, "inspect_audio", soundtrack)
    ing = Ingested("video", "soundtrack-only.mp4", load_media(b"original", "soundtrack-only.mp4"))
    report = asyncio.run(pipeline.analyse(ing, "", "general"))
    assert events == ["original", "audio"]
    assert report["providers"][0]["summary"] == "Original inspected"
    assert report["audio_inspection"]["status"] == "present"
    assert report["meta"]["frame_count"] == 0


def test_deep_pass_carries_only_matching_original_audio_observations():
    original = b"original"
    digest = hashlib.sha256(original).hexdigest()
    prior = {"audio_assessment": {"input_sha256": digest, "content_type": "music"},
             "suno_check": {"input_sha256": digest, "verdict": "verified_suno"}}
    report = {}
    pipeline._carry_audio_observations(prior, report, original)
    assert report == prior
    other_report = {}
    pipeline._carry_audio_observations(prior, other_report, b"different original")
    assert other_report == {}


def test_original_audio_provenance_precedes_transcription(monkeypatch):
    events = []
    row = ProviderResult("c2pa", "C2PA", Kind.PROVENANCE, Status.OK, summary="Original inspected")
    def original(media):
        assert media.raw_bytes == b"original" and media.mime == "audio/mpeg"
        events.append("original")
        return row
    async def transcribe(*_):
        events.append("transcript")
        return None, ["No transcript"]
    monkeypatch.setattr(c2pa_provider, "run", original)
    monkeypatch.setattr(pipeline, "transcribe", transcribe)
    report = asyncio.run(pipeline._analyse_audio(ingest(b"original", "test.mp3"), "", "general"))
    assert events == ["original", "transcript"]
    assert report["providers"][0]["summary"] == "Original inspected"


@pytest.mark.parametrize("endpoint,field", [("/api/audio/assess", "consent_to_google"), ("/api/suno/check", "consent_to_suno")])
def test_optional_endpoints_require_consent_and_matching_hash(monkeypatch, endpoint, field):
    monkeypatch.setattr(main, "settings", replace(main.settings, google_key="fixture", disable_google=False))
    def forbidden(*_):
        raise AssertionError("No provider call may occur")
    monkeypatch.setattr(main, "_original_provenance", forbidden)
    client = TestClient(main.app)
    data = {"expected_sha256": "b" * 64, "track_index": "0"}
    files = {"file": ("test.mp4", b"raw", "video/mp4")}
    assert client.post(endpoint, data=data, files=files).status_code == 403
    data[field] = "true"
    assert client.post(endpoint, data=data, files=files).status_code == 400


def test_optional_audio_orders_provenance_before_derivative(monkeypatch):
    monkeypatch.setattr(main, "settings", replace(main.settings, google_key="fixture", disable_google=False))
    events = []
    async def original(*_):
        events.append("original")
        return {"scope": "original only"}
    def excerpt(*_):
        events.append("excerpt")
        return b"pcm", 1.0
    async def assessed(*_):
        events.append("external")
        return {"status": "ok", "synthetic_audio_assessed": False}
    monkeypatch.setattr(main, "_original_provenance", original)
    monkeypatch.setattr(main, "extract_excerpt", excerpt)
    monkeypatch.setattr(audio_assessment, "assess", assessed)
    response = TestClient(main.app).post("/api/audio/assess",
        files={"file": ("file.mp4", b"raw", "video/mp4")},
        data={"expected_sha256": hashlib.sha256(b"raw").hexdigest(), "track_index": "0", "consent_to_google": "true"})
    assert response.status_code == 200 and events == ["original", "excerpt", "external"]


def test_audio_model_returns_content_only_and_ignores_injected_scores(monkeypatch):
    monkeypatch.setattr(audio_assessment, "settings", replace(audio_assessment.settings, google_key="fixture", disable_google=False))
    calls = []
    def handler(request):
        calls.append(request)
        assert request.url.host == "generativelanguage.googleapis.com"
        data = {"content_type": "music", "observations": ["Instrumental music"],
                "transcript_excerpt": "", "limitations": [], "synthetic_rating": 99}
        return httpx.Response(200, json={"candidates": [{"content": {"parts": [{"text": json.dumps(data)}]}}]})
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            return await audio_assessment.assess(client, b"pcm", "a" * 64, 1, 0, 20)
    result = asyncio.run(run())
    assert result["status"] == "ok" and result["content_type"] == "music"
    assert "synthetic_rating" not in result and result["synthetic_audio_assessed"] is False
    assert len(calls) == 1


@pytest.mark.parametrize("status,answer,expected", [
    (200, {"verdict": "verified_suno"}, "ok"),
    (200, {"verdict": "no_suno_provenance"}, "ok"),
    (200, {"verdict": "new_unsupported_verdict"}, "error"),
    (200, {"verdict": ["verified_suno"]}, "error"),
    (415, {"error": "unsupported"}, "error"),
    (429, {}, "error"), (302, {}, "error"),
])
def test_suno_vendor_contract_has_no_retry_or_negative_on_failure(status, answer, expected):
    calls = []
    def handler(request):
        calls.append(request)
        assert str(request.url) == suno_credentials.ENDPOINT
        assert b"original.mp4" in request.content and b"private-filename" not in request.content
        return httpx.Response(status, json=answer, headers={"location": "https://unrelated.example/"})
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler), follow_redirects=False) as client:
            return await suno_credentials.check(client, b"original bytes", "video/mp4", ".mp4")
    result = asyncio.run(run())
    assert result["status"] == expected and len(calls) == 1
    assert result["input_sha256"] == hashlib.sha256(b"original bytes").hexdigest()
    if expected == "error":
        assert result["verdict"] is None


def test_exports_keep_audio_separate_and_reject_malformed_records():
    report = {"meta": {"tool": "SDA Vision", "filename": "video.mp4", "version": "test",
                       "kind": "video", "generated_at": "2026-09-24", "models": [], "notes": []},
              "consensus": aggregate([], "", "general"), "providers": [], "frames": [],
              "audio_inspection": inspection(),
              "audio_assessment": {"status": "ok", "input_sha256": "a" * 64,
                  "excerpt_sha256": "b" * 64, "provider": "Google Gemini", "model": "fixture",
                  "prompt_version": "soundtrack-content-v1", "checked_at": "2026-09-24",
                  "track_index": 1, "start_seconds": 0, "duration_seconds": 20,
                  "consent_to_google": True, "synthetic_audio_assessed": False,
                  "content_type": "music", "observations": ["Instrumental music"],
                  "transcript_excerpt": "", "limitations": ["Excerpt only"]},
              "suno_check": {"status": "ok", "input_sha256": "a" * 64,
                  "provider": "Suno Credentials", "endpoint": suno_credentials.ENDPOINT,
                  "checked_at": "2026-09-24", "consent_to_suno": True, "submitted_original": True,
                  "verdict": "verified_suno", "http_status": 200, "response_sha256": "c" * 64,
                  "response": {"verdict": "verified_suno"}, "note": "Vendor result, not acoustic detection"}}
    before = copy.deepcopy(report["consensus"])
    validated = validate_report(report)
    assert validated["suno_check"]["response"] == {"verdict": "verified_suno"}
    exports = [build_markdown(validated), build_csv(validated)]
    deck = Presentation(io.BytesIO(build_pptx(validated)))
    exports.append(" ".join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
                   + " " + " ".join(cell.text for slide in deck.slides for shape in slide.shapes if shape.has_table
                                      for row in shape.table.rows for cell in row.cells))
    from pypdf import PdfReader
    pdf = PdfReader(io.BytesIO(build_pdf(validated))).pages
    exports.append(" ".join(page.extract_text() for page in pdf))
    for exported in exports:
        text = " ".join(exported.split())
        assert "Local soundtrack inspection" in text
        assert "Soundtrack-origin detection and continuous-motion analysis are outside these assessments" in text
        assert "Instrumental music" in text and "verified_suno" in text
        assert "soundtrack-content-v1" in text and "not acoustic detection" in text
    assert report["consensus"] == before
    report["audio_inspection"]["tracks"] = "bad"
    with pytest.raises(Exception) as error:
        validate_report(report)
    assert error.value.status_code == 400
