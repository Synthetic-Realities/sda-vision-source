"""Copy regression checks; provider calls are unnecessary for editorial work."""
import ast
import copy
import hashlib
import io
from pathlib import Path

from pypdf import PdfReader
import pytest
from pptx import Presentation

from app.aggregate import aggregate
from app.export import build_csv, build_markdown, build_pdf, build_pptx, _score_basis
from app.schema import Kind, ProviderResult, Status, Verdict, verdict_from_rating

ROOT = Path(__file__).resolve().parents[1]


@pytest.mark.parametrize("path,digest", [
    ("app/providers/prompt.py", "5215591a3c3f9769e8ce1f68c2789a28ca2e2cbf07c9c3322389a5d4a22135da"),
    ("app/providers/transcribe.py", "5f07cba6b04fd1c23fa21ed474758aa9f24d2cdecbd3b6cd8c087b8184095da2"),
    ("app/annotation/prompts.py", "d87d1a93ee188492b59b631f08af65e40eaea5d444be02e85b9f0e35512ebe8a"),
])
def test_editorial_revision_preserves_prompt_sources(path, digest):
    # A later intentional method change must review and update these baselines.
    assert hashlib.sha256((ROOT / path).read_bytes()).hexdigest() == digest


def test_audio_and_copied_second_opinion_prompts_are_unchanged():
    module = ast.parse((ROOT / "app/providers/audio_assessment.py").read_text())
    prompt = next(ast.literal_eval(n.value) for n in module.body
                  if isinstance(n, ast.Assign) and any(isinstance(t, ast.Name) and t.id == "PROMPT" for t in n.targets))
    assert hashlib.sha256(prompt.encode()).hexdigest() == "e3bfb7864951dcd538dccce64012437371aca57b6774a12a72815408a5783004"
    copied = (ROOT / "frontend/src/lib.ts").read_text().split("export function secondOpinionPrompt", 1)[1]
    assert hashlib.sha256(copied.encode()).hexdigest() == "a71a0cce315e82761bafe550834f08a213a3810e4cbec4cdad925376ca6d2506"


@pytest.mark.parametrize("ratings,verdict,phrase", [
    ([], "inconclusive", "No vision model returned a usable result"),
    ([95], "inconclusive", "A corroborated panel result is unavailable"),
    ([95, 10], "inconclusive", "other available models returned different assessments"),
    ([20, 25], "authentic_likely", "2 models lean authentic"),
    ([90, 95], "synthetic_likely", "2 models agree"),
    ([45, 50], "partially_synthetic", "ambiguous or possibly edited"),
])
def test_outcome_copy_keeps_actual_result_and_count(ratings, verdict, phrase):
    providers = [ProviderResult(str(i), f"Fixture {i}", Kind.VISION, Status.OK,
                                rating=rating, verdict=verdict_from_rating(rating))
                 for i, rating in enumerate(ratings)]
    result = aggregate(providers, "", "general")
    assert result["overall_verdict"] == verdict
    assert phrase in result["explanation"]
    assert result["disclaimer"] == "Research assessment. Review alongside source information and context."
    if verdict == "inconclusive":
        assert result["headline"] == "Inconclusive"


def test_unknown_calculation_is_still_not_a_probability():
    assert _score_basis({"consensus": {"overall_rating": 55}}) == (
        "Calculation method unavailable in this report. Not a calibrated probability.")


def test_new_exports_preserve_historical_findings_and_notes():
    evidence = "Historical evidence: not a new result."
    note = "Original researcher note: never silently rewrite this."
    provider = ProviderResult("fixture", "Fixture model", Kind.VISION, Status.OK,
                              rating=50, verdict=Verdict.INCONCLUSIVE,
                              summary=evidence, evidence=[evidence]).to_dict()
    report = {
        "meta": {"tool": "SDA Vision", "filename": "fixture.png", "kind": "image",
                 "version": "fixture", "models": ["fixture"], "generated_at": "2026-09-24",
                 "notes": [], "frame_count": 1},
        "providers": [provider], "frames": [],
        "consensus": aggregate([ProviderResult.from_dict(provider)], "", "general"),
        "second_opinion_gemini": note,
    }
    report["consensus"].update(explanation="Historical explanation: still verify.",
                               disclaimer="Historical disclaimer: not a determination of fact.")
    original = copy.deepcopy(report)
    md, csv = build_markdown(report), build_csv(report)
    pdf = PdfReader(io.BytesIO(build_pdf(report))).pages
    deck = Presentation(io.BytesIO(build_pptx(report)))
    pdf_text = " ".join(page.extract_text() for page in pdf)
    deck_text = " ".join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    # Provider evidence is retained in the deck's speaker notes by the existing exporter.
    deck_text += " " + " ".join(slide.notes_slide.notes_text_frame.text for slide in deck.slides)
    for text in (md, csv, pdf_text, deck_text):
        flat = " ".join(text.split())
        assert evidence in flat
        assert note in flat
        assert "Research assessment." in flat
        assert "calibrated probabilit" in flat.lower()
    assert "Additional findings" in md and "How this assessment was reached" in md
    assert "Additional findings" in pdf_text and "How this assessment was reached" in pdf_text
    for text in (md, pdf_text, deck_text):
        assert "Historical explanation: still verify." in " ".join(text.split())
    assert report == original
    assert "not_decisive" in csv  # Machine field values retain their contract.
