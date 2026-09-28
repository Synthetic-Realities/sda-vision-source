"""Report layout regressions with synthetic records and no provider calls."""
import copy
import base64
import io

from pypdf import PdfReader
from pptx import Presentation
from PIL import Image

from app.aggregate import aggregate
from app.export import build_pdf, build_pptx, build_batch_pdf, build_batch_pptx


def fixture():
    return {
        "meta": {"tool": "SDA Vision", "filename": "layout-fixture.png", "kind": "image",
                 "generated_at": "2026-09-25", "models": ["fixture"], "notes": [], "frame_count": 1},
        "consensus": {**aggregate([], "", "general"), "visible_text": "Transcript text. " * 150 + "TEXT END MARKER"},
        "frames": [],
        "providers": [{"id": "fixture", "name": "Fixture model", "kind": "vision", "status": "ok",
                       "verdict": "inconclusive", "rating": None, "summary": "Café & <scope>.",
                       "evidence": [f"Evidence item {i}." for i in range(1, 10)]}],
        "second_opinion_gemini": "Researcher-supplied observation. " * 190 + "NOTE END MARKER",
    }


def test_tables_preserve_all_evidence_and_paginate_without_mutation():
    report = fixture()
    original = copy.deepcopy(report)
    pdf = PdfReader(io.BytesIO(build_pdf(report))).pages
    deck = Presentation(io.BytesIO(build_pptx(report)))
    pdf_text = " ".join(page.extract_text() for page in pdf)
    pptx_text = " ".join(cell.text for slide in deck.slides for shape in slide.shapes if shape.has_table
                         for row in shape.table.rows for cell in row.cells)
    for text in (pdf_text, pptx_text):
        flat = " ".join(text.split())
        for marker in ["Evidence item 9.", "TEXT END MARKER", "NOTE END MARKER", "Café & <scope>."]:
            assert marker in flat
    assert report == original
    assert len(pdf) > 2 and len(deck.slides) > 3
    # Check visible glyph bounds with an independent renderer, preserving the
    # off-page-content assertion after removing the old PDF library.
    import pypdfium2 as pdfium
    from contextlib import closing
    with pdfium.PdfDocument(build_pdf(report)) as document:
        for index in range(len(document)):
            with closing(document[index]) as page, closing(page.get_textpage()) as textpage:
                width, height = page.get_size()
                for i in range(textpage.count_chars()):
                    if not textpage.get_text_range(i, 1).strip():
                        continue
                    left, bottom, right, top = textpage.get_charbox(i)
                    assert 0 <= left <= right <= width
                    assert 0 <= bottom <= top <= height
    for slide in deck.slides:
        assert any(shape.has_table for shape in slide.shapes)
        for shape in slide.shapes:
            assert shape.left + shape.width <= deck.slide_width
            assert shape.top + shape.height <= deck.slide_height


def test_batch_exports_retain_every_report(monkeypatch):
    monkeypatch.setattr("app.export._batch_graph", lambda reports: None)
    first, second = fixture(), fixture()
    second["meta"]["filename"] = "SECOND REPORT"
    items = [{"report": first}, {"report": second}]
    pdf = PdfReader(io.BytesIO(build_batch_pdf(items))).pages
    deck = Presentation(io.BytesIO(build_batch_pptx(items)))
    assert "SECOND REPORT" in " ".join(page.extract_text() for page in pdf)
    assert "SECOND REPORT" in " ".join(cell.text for slide in deck.slides for shape in slide.shapes
                                      if shape.has_table for row in shape.table.rows for cell in row.cells)


def test_full_research_exports_embed_supplied_media_preview():
    image = io.BytesIO()
    Image.new("RGB", (600, 300), "#15878a").save(image, format="PNG")
    encoded = "data:image/png;base64," + base64.b64encode(image.getvalue()).decode()
    report = fixture()
    pdf = PdfReader(io.BytesIO(build_pdf(report, encoded))).pages
    assert pdf[0].images
    deck = Presentation(io.BytesIO(build_pptx(report, encoded)))
    pictures = [shape for shape in deck.slides[0].shapes if shape.shape_type == 13]
    assert len(pictures) == 1
    assert abs(pictures[0].width / pictures[0].height - 2) < .01
