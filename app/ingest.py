"""Turn any uploaded file into analysable units.

Visual documents (PDF, PPTX, video) are expanded into a capped set of sampled
frames (images) that the vision + forensic pipeline runs on per frame. The
original bytes are kept for whole-file providers (C2PA reads PDF/MP4 manifests
directly). Plain transcripts become a text unit for content analysis.
"""
from __future__ import annotations

import io
from contextlib import closing
import os
import subprocess
import tempfile
import threading
from dataclasses import dataclass, field
from pathlib import Path

from PIL import Image

from .config import settings
from .media import Media, load_media

# Render quality for document pages / slides.
_PAGE_DPI = 110
_JPEG_Q = 85

# PDFium is process-global and not thread-safe, including across documents.
_PDFIUM_LOCK = threading.Lock()
PDF_INPUT_REVISION = "pdf-input-2026-09-26.1"

_IMAGE_EXT = {".jpg", ".jpeg", ".png", ".webp", ".gif", ".heic", ".heif", ".bmp", ".tif", ".tiff"}
_TEXT_EXT = {".txt", ".srt", ".vtt", ".md"}
_VIDEO_EXT = {".mp4", ".mov", ".webm", ".m4v", ".avi"}
_AUDIO_EXT = {".m4a", ".mp3", ".wav", ".aac", ".flac"}


@dataclass
class Frame:
    media: Media
    label: str  # e.g. "image", "page 2", "slide 3", "00:48"


@dataclass
class Ingested:
    kind: str                       # image|pdf|pptx|video|text|audio|unsupported
    filename: str
    original: Media                 # whole-file bytes (for C2PA etc.)
    frames: list[Frame] = field(default_factory=list)
    text: str | None = None         # transcript / extracted document text
    notes: list[str] = field(default_factory=list)
    # How many frames QUALIFIED before the frame budget was applied. For finite
    # sources (PDF figures/pages, PPTX images) this is the full count, so the UI
    # can offer a deep pass over the remainder; video is continuous, so it just
    # mirrors the sampled count.
    frames_found: int = 0


# C2PA reads manifests straight from these containers, so the original needs the
# right MIME (the byte-sniffer only knows image formats).
_EXT_MIME = {
    ".pdf": "application/pdf",
    ".mp4": "video/mp4", ".m4v": "video/mp4", ".mov": "video/quicktime",
    ".webm": "video/webm", ".avi": "video/x-msvideo",
    ".m4a": "audio/mp4", ".mp3": "audio/mpeg", ".wav": "audio/wav",
    ".aac": "audio/aac", ".flac": "audio/flac",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}


def ingest(raw: bytes, filename: str, mime: str | None = None,
           max_frames: int | None = None) -> Ingested:
    """max_frames overrides the global frame budget for this call only:
    None = use settings.MAX_FRAMES, 0 = no cap (every qualifying frame), N = N.
    Used by the deep pass, which needs the full frame set."""
    ext = Path(filename).suffix.lower()
    original = load_media(raw, filename, mime)
    if ext in _EXT_MIME:
        original.mime = _EXT_MIME[ext]
    cap = settings.max_frames if max_frames is None else max_frames

    if ext in _IMAGE_EXT or (original.width and ext not in _TEXT_EXT | _VIDEO_EXT | _AUDIO_EXT | {".pdf", ".pptx"}):
        return Ingested("image", filename, original, frames=[Frame(original, "image")],
                        frames_found=1)
    if ext == ".pdf":
        return _ingest_pdf(raw, filename, original, cap)
    if ext == ".pptx":
        return _ingest_pptx(raw, filename, original, cap)
    if ext in _VIDEO_EXT:
        return _ingest_video(raw, filename, original, cap)
    if ext in _TEXT_EXT:
        return _ingest_text(raw, filename, original)
    if ext in _AUDIO_EXT:
        return Ingested("audio", filename, original)
    return Ingested("unsupported", filename, original, notes=[f"Unsupported file type: {ext or 'unknown'}."])


def _frame_from_image_bytes(png_or_jpg: bytes, name: str, label: str) -> Frame:
    return Frame(load_media(png_or_jpg, name), label)


def _to_jpeg(img: Image.Image) -> bytes:
    buf = io.BytesIO()
    img.convert("RGB").save(buf, format="JPEG", quality=_JPEG_Q)
    return buf.getvalue()


def _even_indices(total: int, cap: int) -> list[int]:
    if cap <= 0 or total <= cap:  # cap 0 = no budget: take everything
        return list(range(total))
    if cap == 1:  # a budget of one means the first page/slide, not a midpoint
        return [0]
    # Evenly spaced sample across the whole range.
    return [round(i * (total - 1) / (cap - 1)) for i in range(cap)]


def _ingest_pdf(raw: bytes, filename: str, original: Media, cap: int) -> Ingested:
    from pypdf import PdfReader

    notes: list[str] = [
        f"PDF preparation: {PDF_INPUT_REVISION} (pypdf text/figures; PDFium page rendering)."
    ]
    try:
        doc = PdfReader(io.BytesIO(raw))
        # Empty-password PDFs may be read; a password prompt is outside this workflow.
        if doc.is_encrypted and not doc.decrypt(""):
            return Ingested("pdf", filename, original,
                            notes=notes + ["Could not open PDF: a password is required."])
        page_count = len(doc.pages)
    except Exception as exc:
        return Ingested("pdf", filename, original, notes=notes + [f"Could not open PDF: {exc}"])

    text_chunks = []
    for i, page in enumerate(doc.pages):
        try:
            text = page.extract_text().strip()
            if text:
                text_chunks.append(text)
        except Exception:
            notes.append(f"Selectable text could not be extracted from page {i + 1}.")

    frames: list[Frame] = []
    found = 0
    mode = settings.pdf_mode
    if mode in ("auto", "figures"):
        frames, found = _extract_pdf_figures(doc, filename, notes, cap)
    if not frames and mode != "figures":
        frames, found = _render_pdf_pages(raw, page_count, filename, notes, cap)
    elif not frames and mode == "figures":
        notes.append("No embedded figures passed the size/shape filter.")
    if not text_chunks and frames:
        notes.append("No selectable text layer - vision models read the page/figure images.")
    return Ingested("pdf", filename, original, frames=frames,
                    text="\n\n".join(text_chunks) or None, notes=notes, frames_found=found)


def _render_pdf_pages(raw: bytes, page_count: int, filename: str,
                      notes: list[str], cap: int) -> tuple[list[Frame], int]:
    import pypdfium2 as pdfium

    frames = []
    pages = _even_indices(page_count, cap)
    # Hold the mutex through construction, rendering and cleanup of native handles.
    with _PDFIUM_LOCK:
        try:
            with pdfium.PdfDocument(raw) as doc:
                for i in pages:
                    try:
                        with closing(doc[i]) as page:
                            with closing(page.render(scale=_PAGE_DPI / 72)) as bitmap:
                                with bitmap.to_pil() as img:
                                    blob = _to_jpeg(img)
                        frames.append(_frame_from_image_bytes(blob, f"{filename}-p{i+1}.jpg", f"page {i + 1}"))
                    except Exception:
                        notes.append(f"Page {i + 1} could not be rendered.")
        except Exception:
            notes.append("PDF page rendering failed.")
    notes.append(f"Rendered {len(frames)} of {page_count} page(s) at {_PAGE_DPI} dpi.")
    return frames, page_count


def _extract_pdf_figures(doc, filename: str, notes: list[str], cap: int) -> tuple[list[Frame], int]:
    """Extract raster figures; preserve size/shape filters and stable full-list labels.

    Deduplicate indirect image objects across pages, not distinct objects whose
    pixels happen to be identical. Inline images have page-local identities.
    """
    min_px = settings.min_figure_px
    seen: set[tuple] = set()
    figs: list[tuple[int, bytes]] = []
    for pno, page in enumerate(doc.pages):
        try:
            keys = page.images.keys()
        except Exception:
            notes.append(f"Embedded figures could not be listed on page {pno + 1}.")
            continue
        for key in keys:
            try:
                item = page.images[key]
                ref = item.indirect_reference
                identity = ("object", ref.idnum, ref.generation) if ref else ("inline", pno, str(key))
                if identity in seen:
                    continue
                seen.add(identity)
                with Image.open(io.BytesIO(item.data)) as img:
                    w, h = img.size
                if w < min_px or h < min_px or max(w, h) / max(1, min(w, h)) > 5:
                    continue
                figs.append((pno + 1, item.data))
            except Exception:
                notes.append(f"An embedded figure on page {pno + 1} could not be extracted.")

    seen_pages: dict[int, int] = {}
    labels: list[str] = []
    for pno, _ in figs:
        seen_pages[pno] = seen_pages.get(pno, 0) + 1
        labels.append(f"figure p{pno}" if seen_pages[pno] == 1
                      else f"figure p{pno}.{seen_pages[pno]}")
    frames = []
    for j in _even_indices(len(figs), cap):
        pno, blob = figs[j]
        try:
            with Image.open(io.BytesIO(blob)) as img:
                frame = _frame_from_image_bytes(_to_jpeg(img), f"{filename}-fig{j}.jpg", labels[j])
            frames.append(frame)
        except Exception:
            notes.append(f"Figure {labels[j]} could not be prepared.")
    if frames:
        notes.append(f"Extracted {len(frames)} embedded figure(s) of {len(figs)} found (figure mode).")
    return frames, len(figs)


def _ingest_pptx(raw: bytes, filename: str, original: Media, cap: int) -> Ingested:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    notes: list[str] = []
    try:
        prs = Presentation(io.BytesIO(raw))
    except Exception as exc:
        return Ingested("pptx", filename, original, notes=[f"Could not open PPTX: {exc}"])

    slide_imgs: list[tuple[int, bytes]] = []
    text_chunks: list[str] = []
    for idx, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                text_chunks.append(shape.text_frame.text.strip())
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    slide_imgs.append((idx, shape.image.blob))
                except Exception:
                    pass

    # Sample across the available slide images.
    picks = _even_indices(len(slide_imgs), cap)
    frames = []
    for j in picks:
        sidx, blob = slide_imgs[j]
        try:
            img = Image.open(io.BytesIO(blob))
            frames.append(_frame_from_image_bytes(_to_jpeg(img), f"{filename}-s{sidx}.jpg", f"slide {sidx}"))
        except Exception:
            continue

    notes.append(f"{len(prs.slides)} slides; sampled {len(frames)} of {len(slide_imgs)} embedded slide image(s).")
    if not frames:
        notes.append("No embedded slide images found; only slide text was extracted "
                     "(full visual rendering would need LibreOffice).")
    return Ingested("pptx", filename, original, frames=frames,
                    text="\n\n".join(text_chunks) or None, notes=notes,
                    frames_found=len(slide_imgs))


def _ingest_video(raw: bytes, filename: str, original: Media, cap: int) -> Ingested:
    if not _have("ffmpeg"):
        return Ingested("video", filename, original,
                        notes=["ffmpeg not found on PATH - install it to enable video frame sampling."])
    notes: list[str] = []
    with tempfile.TemporaryDirectory() as td:
        src = os.path.join(td, "in" + Path(filename).suffix)
        with open(src, "wb") as f:
            f.write(raw)
        duration = _video_duration(src)
        n = cap if cap > 0 else settings.max_frames
        # Sample evenly across the timeline, skipping the very start/end.
        if duration and duration > 1:
            stamps = [duration * (k + 1) / (n + 1) for k in range(n)]
        else:
            stamps = [0.0]
        frames = []
        for k, ts in enumerate(stamps):
            out = os.path.join(td, f"f{k:02d}.jpg")
            r = subprocess.run(
                ["ffmpeg", "-v", "quiet", "-ss", f"{ts:.2f}", "-i", src,
                 "-frames:v", "1", "-q:v", "3", out],
                capture_output=True)
            if r.returncode == 0 and os.path.exists(out):
                with open(out, "rb") as fh:
                    frames.append(_frame_from_image_bytes(fh.read(), f"{filename}-f{k}.jpg", _hms(ts)))
        notes.append(f"Duration {duration:.0f}s; sampled {len(frames)} frame(s)." if duration
                     else f"Sampled {len(frames)} frame(s).")
    return Ingested("video", filename, original, frames=frames, notes=notes,
                    frames_found=len(frames))


def _ingest_text(raw: bytes, filename: str, original: Media) -> Ingested:
    text = raw.decode("utf-8", errors="replace").strip()
    notes = [f"{len(text)} characters of transcript/text."]
    if not text:
        notes = ["File is empty."]
    return Ingested("text", filename, original, text=text or None, notes=notes)


def _have(binary: str) -> bool:
    from shutil import which
    return which(binary) is not None


def _video_duration(path: str) -> float:
    try:
        out = subprocess.run(
            ["ffprobe", "-v", "quiet", "-show_entries", "format=duration",
             "-of", "csv=p=0", path], capture_output=True, text=True).stdout.strip()
        return float(out)
    except Exception:
        return 0.0


def _hms(seconds: float) -> str:
    s = int(seconds)
    return f"{s // 60:02d}:{s % 60:02d}"
