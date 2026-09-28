"""Shared table layout for research PDF and PowerPoint exports."""
from __future__ import annotations

import io
from pathlib import Path
from xml.sax.saxutils import escape


def report_font():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    name = "SDA-NotoSans"
    if name not in pdfmetrics.getRegisteredFontNames():
        path = Path(__file__).resolve().parents[1] / "frontend/src/assets/fonts/NotoSans-Regular.ttf"
        pdfmetrics.registerFont(TTFont(name, str(path)))
    return name


def _slide_lines(text, width, size):
    # Explicit line breaks and long identifiers both need bounded table cells.
    from reportlab.pdfbase.pdfmetrics import stringWidth
    out = []
    for line in str(text).split("\n"):
        current = ""
        for char in line:
            if current and stringWidth(current + char, report_font(), size) > width:
                split = current.rfind(" ")
                if split > 0:
                    out.append(current[:split])
                    current = current[split + 1:] + char
                else:
                    out.append(current)
                    current = char
            else:
                current += char
        out.append(current)
    return out or [""]


def add_pptx_sections(prs, sections, filename, scope, image=None):
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    def text(slide, value, x, y, w, h, size=12, color="21243A", bold=False):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = 0
        tf.margin_top = tf.margin_bottom = 0
        tf.text = value
        for p in tf.paragraphs:
            p.font.name = "Arial"
            p.font.size = Pt(size)
            p.font.bold = bold
            p.font.color.rgb = RGBColor.from_string(color)
            p.space_after = Pt(0)
            p.line_spacing = 1.15
        return shape

    def page(title, notes=""):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        text(slide, "SDA VISION / RESEARCH REPORT", .4, .24, 9.2, .25, 11, "C5006A", True)
        text(slide, title, .4, .72, 9.2, .67, 22, bold=True)
        text(slide, scope, .4, 6.88, 8.75, .4, 8, "545B6D")
        text(slide, str(len(prs.slides)), 9.18, 6.91, .4, .25, 9, "545B6D")
        slide.notes_slide.notes_text_frame.text = filename + "\n" + notes
        return slide

    def table(slide, values, widths, y, heights, header=False, shaded=False):
        t = slide.shapes.add_table(len(values), len(widths), Inches(.4), Inches(y),
                                   Inches(sum(widths)), Inches(sum(heights))).table
        for i, width in enumerate(widths):
            t.columns[i].width = Inches(width)
        for i, row in enumerate(values):
            t.rows[i].height = Inches(heights[i])
            for j, value in enumerate(row):
                cell = t.cell(i, j)
                cell.margin_left = cell.margin_right = Inches(.09)
                cell.margin_top = cell.margin_bottom = Inches(.07)
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string("16787C" if header else "F1F7F7" if shaded else "FFFFFF")
                cell.text = str(value)
                cell.text_frame.word_wrap = False
                for p in cell.text_frame.paragraphs:
                    p.font.name = "Arial"
                    p.font.size = Pt(12)
                    p.font.bold = header
                    p.font.color.rgb = RGBColor.from_string("FFFFFF" if header else "21243A")
                    p.line_spacing = Pt(16)
                    p.space_after = Pt(0)

    if image:
        from PIL import Image
        with Image.open(io.BytesIO(image)) as im:
            w, h = im.size
        slide = page("Analysed item", filename)
        iw, ih = min(8.5, 4.5 * w / h), min(4.5, 8.5 * h / w)
        slide.shapes.add_picture(io.BytesIO(image), Inches((10 - iw) / 2), Inches(1.48), width=Inches(iw), height=Inches(ih))
        lines = _slide_lines(filename, 9.2 * 72, 11)
        text(slide, "\n".join(lines[:3]), .4, 6.05, 9.2, .65, 11)

    for section in sections:
        widths = [value * 9.2 for value in section["widths"]]
        notes = section["title"] + "\n" + "\n".join(" | ".join(map(str, row)) for row in section["rows"])
        part = 0

        def start_page():
            nonlocal part
            part += 1
            slide = page(section["title"] + (" (continued)" if part > 1 else ""), notes)
            table(slide, [section["headers"]], widths, 1.48, [.38], header=True)
            return slide

        slide, y = start_page(), 1.86
        for index, row in enumerate(section["rows"]):
            cells = [_slide_lines(value, width * 72 - 16, 12) for value, width in zip(row, widths)]
            count, start = max(map(len, cells)), 0
            while start < count:
                available = int((6.65 - y - .16) / (16 / 72))
                if available < 2:
                    slide, y = start_page(), 1.86
                    available = 20
                take = min(available, count - start)
                values = ["\n".join(cell[start:start + take]) or ("Continued" if i == 0 else "")
                          for i, cell in enumerate(cells)]
                height = take * 16 / 72 + .16
                table(slide, [values], widths, y, [height], shaded=index % 2 == 0)
                y += height
                start += take


def pdf_sections(sections, filename, tool, image=None):
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Image as RLImage, Paragraph, Spacer, Table, TableStyle

    font = report_font()
    normal = ParagraphStyle("ReportCell", fontName=font, fontSize=9, leading=13, splitLongWords=True)
    head = ParagraphStyle("ReportHeader", parent=normal, textColor=colors.white)
    title = ParagraphStyle("ReportTitle", parent=normal, fontName="Helvetica-Bold", fontSize=23, leading=28,
                           textColor=colors.HexColor("#C5006A"), spaceAfter=12)
    section_style = ParagraphStyle("ReportSection", parent=normal, fontName="Helvetica-Bold", fontSize=14, leading=18,
                                   textColor=colors.HexColor("#126E73"), spaceBefore=17, spaceAfter=9, keepWithNext=True)
    width = 510

    def para(value, style=normal):
        return Paragraph(escape(str(value)).replace("\n", "<br/>"), style)

    flow = [para(f"{tool} / Research report", title), para(filename), Spacer(1, 14)]
    if image:
        from PIL import Image
        with Image.open(io.BytesIO(image)) as im:
            w, h = im.size
        iw, ih = min(280, 160 * w / h), min(160, 280 * h / w)
        flow += [RLImage(io.BytesIO(image), width=iw, height=ih), Spacer(1, 10)]
    for section in sections:
        flow.append(para(section["title"], section_style))
        data = [[para(value, head) for value in section["headers"]]]
        data += [[para(value) for value in row] for row in section["rows"]]
        table = Table(data, colWidths=[v * width for v in section["widths"]], repeatRows=1, splitByRow=1, splitInRow=1, hAlign="LEFT")
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#16787C")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#F1F7F7"), colors.white]),
            ("GRID", (0, 0), (-1, -1), .4, colors.HexColor("#CFDBDF")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ]))
        flow.append(table)
    return flow


def pdf_page(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColorRGB(.33, .36, .43)
    canvas.drawString(43, 26, "SDA Vision | Research report | Ratings are not calibrated probabilities.")
    canvas.drawRightString(552, 26, str(doc.page))
    canvas.restoreState()
